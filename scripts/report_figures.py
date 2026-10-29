"""Generate the figures used by the major-project report into docs/report_assets/.

    python scripts/report_figures.py            # all figures
    python scripts/report_figures.py diagrams   # only the design diagrams (no models needed)

Diagrams are drawn with matplotlib. Result figures come from docs/paper_assets (made by
scripts/paper_experiments.py), the executed training notebook, and a run of the real pipeline.
"""

import base64
import json
import shutil
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, FancyArrowPatch, FancyBboxPatch, Rectangle

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "report_assets"
OUT.mkdir(parents=True, exist_ok=True)

plt.rcParams.update({"font.family": "serif", "font.serif": ["Times New Roman", "DejaVu Serif"],
                     "font.size": 8})
NAVY, TEAL, ORANGE, RED, GRAY = "#1f4e79", "#2a8a8c", "#d98a2b", "#b94a48", "#555555"
LIGHT = "#eef3f8"


def canvas(w, h, xmax=100, ymax=None):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, xmax)
    ax.set_ylim(0, ymax or xmax * h / w)
    ax.axis("off")
    return fig, ax


def box(ax, x, y, w, h, text, fc=LIGHT, ec=NAVY, fs=7.5, bold=False, lw=1.0, round_=1.0):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0.0,rounding_size={round_}",
                                fc=fc, ec=ec, lw=lw))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs,
            weight="bold" if bold else "normal", linespacing=1.25)


def arrow(ax, a, b, text=None, rad=0.0, color="#222", fs=6.5, dy=1.2, dx=0.0, ls="-", head=7):
    ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=head, lw=0.9, color=color,
                                 connectionstyle=f"arc3,rad={rad}", linestyle=ls))
    if text:
        ax.text((a[0] + b[0]) / 2 + dx, (a[1] + b[1]) / 2 + dy, text, ha="center", va="bottom",
                fontsize=fs, color=color)


def save(fig, name):
    fig.savefig(OUT / name, dpi=300, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    print("wrote", name)


# ---------------------------------------------------------------------------------------
def fig_hld():
    """System components and deployment view (what runs where); complements the layered PPT figure."""
    fig, ax = canvas(6.4, 5.0, xmax=100, ymax=78)
    ax.add_patch(Rectangle((0.5, 58), 99, 19.5, fc="none", ec="#999", lw=0.8, ls="--"))
    ax.text(2, 75.2, "Clients", fontsize=7, color=GRAY, style="italic")
    box(ax, 4, 60, 42, 13, "Android phone app (Expo / React Native)\ncamera frames, GPS, compass,\nspeech, OpenStreetMap view", fc="#e7f0f7", fs=7)
    box(ax, 54, 60, 42, 13, "Browser console (HTML / CSS / JS)\nlive feed, alert card, route box,\nvoice test, video upload", fc="#e7f0f7", fs=7)

    box(ax, 4, 48, 92, 6, "Web server (web_app.py)   /frame   /stream   /state   /upload   /places   /location   /route",
        fc="#dfe9d9", ec="#4a7c3a", fs=7, bold=True)
    arrow(ax, (25, 60), (25, 54.4), "JPEG, GPS, heading", dy=-0.8, dx=12)
    arrow(ax, (75, 60), (75, 54.4), "MJPEG stream, JSON", dy=-0.8, dx=12)

    ax.add_patch(Rectangle((0.5, 30), 99, 15, fc="none", ec="#999", lw=0.8, ls="--"))
    ax.text(98.5, 42.4, "Inference engine (Python, PyTorch, CUDA)", fontsize=7, color=GRAY, style="italic", ha="right")
    names = [("Perception", "YOLO11s, hazard\nYOLO11s, ByteTrack", "#e0f0f0", TEAL),
             ("Depth", "Depth Anything V2\nMetric Outdoor Small", "#e0f0f0", TEAL),
             ("Spatial reasoning", "ground plane, scale fix\n1 m x 3 m corridor", "#fbefd9", ORANGE),
             ("Risk and alerts", "TTC, urgency, ranking,\ncooldowns", "#f6dcdb", RED)]
    xs = [8, 29, 50, 71]
    for x, (head, body, fc, ec) in zip(xs, names):
        box(ax, x, 32, 19, 9, f"{head}\n{body}", fc=fc, ec=ec, fs=6.2)
    for x in xs[:-1]:
        arrow(ax, (x + 19.2, 36.5), (x + 21 - 0.2, 36.5))
    arrow(ax, (17, 48), (17, 41.4), "frames", dy=-1.5, dx=4)

    box(ax, 1.5, 7, 28, 14, "Navigation manager\nhaversine, bearing,\nroute steps, off-route", fc="#efe8f6", ec="#6b4c9a", fs=6.5)
    box(ax, 1.5, 0.3, 28, 4.8, "OSM services (OSRM, Nominatim)", fc="#f1f1f1", ec=GRAY, fs=5.8)
    arrow(ax, (3.3, 48), (3.3, 21.4), None)
    ax.text(4.2, 26.5, "GPS,\nroute", fontsize=6)
    arrow(ax, (15.5, 7), (15.5, 5.4), None)
    box(ax, 37, 7, 28, 14, "Annotated frames\n+ state JSON\n(path, alerts, route)", fc="#fff4d6", ec="#a88400", fs=6.5)
    box(ax, 70, 7, 28, 14, "Voice (Windows SAPI,\nphone speech) and stereo\nbeep from obstacle side", fc="#fff4d6", ec="#a88400", fs=6.3)
    arrow(ax, (59.5, 32), (51, 21.4), None)
    arrow(ax, (80.5, 32), (84, 21.4), None)
    save(fig, "fig_components.png")


def stick_figure(ax, x, y, label):
    ax.add_patch(Ellipse((x, y + 7), 3.2, 3.2, fc="white", ec="#222", lw=1))
    ax.plot([x, x], [y + 5.4, y], color="#222", lw=1)
    ax.plot([x - 3, x + 3], [y + 3.8, y + 3.8], color="#222", lw=1)
    ax.plot([x, x - 2.5], [y, y - 4], color="#222", lw=1)
    ax.plot([x, x + 2.5], [y, y - 4], color="#222", lw=1)
    ax.text(x, y - 5.5, label, ha="center", va="top", fontsize=7)


def usecase(ax, x, y, text, w=24, h=7.5, fc="#eef3f8"):
    ax.add_patch(Ellipse((x, y), w, h, fc=fc, ec=NAVY, lw=1))
    ax.text(x, y, text, ha="center", va="center", fontsize=6.8, linespacing=1.15)


def fig_usecase_walk():
    fig, ax = canvas(6.4, 4.8, xmax=100, ymax=76)
    ax.add_patch(Rectangle((24, 1), 52, 73, fc="none", ec=NAVY, lw=1.2))
    ax.text(50, 70.5, "Walking assistance module", ha="center", fontsize=8, weight="bold", color=NAVY)
    stick_figure(ax, 9, 36, "Pedestrian\n(user)")
    stick_figure(ax, 91, 36, "Camera /\nvideo source")
    cases = [(50, 62, "Start / stop monitoring", True, False), (50, 52.5, "View annotated live feed", True, False),
             (50, 43, "Receive obstacle alert\n(distance, direction)", True, True),
             (50, 33.5, "Receive stairs warning\n(up / down)", True, True),
             (50, 24, "Hear directional beep", True, False), (50, 14.5, "Upload video for analysis", True, True),
             (50, 6, "Test voice output", True, False)]
    for x, y, t, user, camera in cases:
        usecase(ax, x, y, t, h=7.2)
        if user:
            ax.plot([12, x - 12], [36, y], color="#555", lw=0.7)
        if camera:
            ax.plot([88, x + 12], [36, y], color="#555", lw=0.7)
    save(fig, "fig_usecase_walk.png")


def fig_usecase_route():
    fig, ax = canvas(6.4, 4.5, xmax=100, ymax=70)
    ax.add_patch(Rectangle((24, 1), 52, 67, fc="none", ec=NAVY, lw=1.2))
    ax.text(50, 64.5, "Route guidance module", ha="center", fontsize=8, weight="bold", color=NAVY)
    stick_figure(ax, 9, 33, "Pedestrian\n(user)")
    stick_figure(ax, 91, 33, "OpenStreetMap\nservices")
    cases = [(50, 56, "Search destination"), (50, 47, "Choose place suggestion"), (50, 38, "Create walking route"),
             (50, 29, "Align phone with the road"), (50, 20, "Receive distance / turn guidance"),
             (50, 12.5, "Off-route warning"), (50, 5.5, "Arrival notification")]
    for x, y, t in cases:
        usecase(ax, x, y, t, h=5.8, w=26)
        ax.plot([12, x - 13], [33, y], color="#555", lw=0.7)
    for y in (56, 38):
        ax.plot([88, 63], [33, y], color="#555", lw=0.7)
    save(fig, "fig_usecase_route.png")


def fig_dfd():
    fig, ax = canvas(6.4, 5.2, xmax=100, ymax=81)
    def entity(x, y, w, h, t):
        ax.add_patch(Rectangle((x, y), w, h, fc="#f4f4f4", ec="#222", lw=1))
        ax.text(x + w / 2, y + h / 2, t, ha="center", va="center", fontsize=7, weight="bold")

    def proc(x, y, t):
        ax.add_patch(Ellipse((x, y), 18, 11.5, fc=LIGHT, ec=NAVY, lw=1))
        ax.text(x, y, t, ha="center", va="center", fontsize=6.4, linespacing=1.15)

    def store(x, y, w, t):
        ax.plot([x, x + w], [y + 4, y + 4], color="#222", lw=1)
        ax.plot([x, x + w], [y, y], color="#222", lw=1)
        ax.text(x + w / 2, y + 2, t, ha="center", va="center", fontsize=6.5)

    def label(x, y, t, ha="center"):
        ax.text(x, y, t, fontsize=6.2, ha=ha, va="bottom")

    entity(1, 62, 16, 9, "Video source /\nphone camera")
    entity(1, 36, 16, 9, "User\n(phone / browser)")
    entity(1, 8, 16, 9, "OpenStreetMap\nservices")
    proc(32, 66.5, "1.0\nDecode\nframe")
    proc(56, 66.5, "2.0\nDetect and\ntrack objects")
    proc(80, 66.5, "3.0\nEstimate\ndepth")
    proc(80, 42, "4.0\nSpatial\nanalysis")
    proc(56, 42, "5.0\nRank and\nannounce alerts")
    proc(32, 42, "7.0\nRender\noutput")
    proc(32, 14, "6.0\nRoute\nguidance")
    arrow(ax, (17, 66.5), (23, 66.5), None); label(20, 67.6, "frame")
    arrow(ax, (41, 66.5), (47, 66.5), None); label(44, 67.6, "image")
    arrow(ax, (65, 66.5), (71, 66.5), None); label(68, 67.6, "boxes")
    arrow(ax, (80, 60.7), (80, 47.8), None); label(81.5, 53.5, "boxes + depth map", ha="left")
    arrow(ax, (71, 42), (65, 42), None); label(68, 43, "path")
    arrow(ax, (47, 42), (41, 42), None); label(44, 43, "alerts")
    arrow(ax, (23, 42), (17, 42), None); label(20, 43, "output")
    arrow(ax, (9, 36), (24, 18.2), None, rad=0.1); label(16.5, 29, "GPS, heading,\ndestination", ha="left")
    arrow(ax, (32, 20), (32, 36.2), None); label(33.5, 27.5, "guidance", ha="left")
    arrow(ax, (17, 19), (24, 16.6), None); label(19.5, 20.2, "route")
    arrow(ax, (24, 12), (17, 12.6), None); label(20.5, 8.6, "query")
    store(46, 76, 40, "D1  Model weights (YOLO, depth)")
    arrow(ax, (56, 76), (56, 72.4), None, ls=":", color=GRAY)
    arrow(ax, (80, 76), (80, 72.4), None, ls=":", color=GRAY)
    store(58, 22, 36, "D2  Session state (in memory)")
    arrow(ax, (56, 36.2), (62, 26.4), None, ls=":", color=GRAY)
    arrow(ax, (41, 14), (58, 23), None, ls=":", color=GRAY)
    store(1, 50, 28, "D3  Uploaded videos")
    arrow(ax, (9, 62), (9, 54.4), None, ls=":", color=GRAY)
    save(fig, "fig_dfd.png")


def fig_sequence():
    names = ["Phone app", "Web server", "Inference\nengine", "Detector", "Depth\nestimator", "Spatial\nanalyzer", "Announcer", "Speaker"]
    n = len(names)
    fig, ax = canvas(6.6, 5.2, xmax=100, ymax=100)
    xs = [6 + i * (88 / (n - 1)) for i in range(n)]
    for x, name in zip(xs, names):
        box(ax, x - 5.2, 92, 10.4, 6.5, name, fs=6.5, bold=True)
        ax.plot([x, x], [92, 8], color="#888", lw=0.7, ls="--")
    msgs = [
        (0, 1, 88, "1  POST /frame (JPEG)"), (1, 2, 83.5, "2  frame(image)"), (2, 3, 79, "3  detect()"),
        (3, 2, 75, "4  detections, track ids"), (2, 4, 71, "5  estimate()"), (4, 2, 67, "6  depth map"),
        (2, 5, 63, "7  analyze(depth, boxes)"), (5, 2, 59, "8  PathStatus"),
        (2, 6, 55, "9  rank_alerts, decide()"), (6, 2, 51, "10  alert (or silence)"),
        (2, 7, 47, "11  say(text), worker thread"), (2, 1, 41, "12  annotated frame + state"),
        (1, 0, 37, "13  JSON (image, state)"),
    ]
    for a, b, y, t in msgs:
        arrow(ax, (xs[a], y), (xs[b], y), t, dy=0.7, fs=6.0, color="#222")
    ax.add_patch(Rectangle((xs[0] - 4, 10), xs[2] - xs[0] + 8, 22, fc="#f6f6f6", ec="#aaa", lw=0.6, zorder=0))
    ax.text(xs[0] - 3, 29.5, "loop: GPS every second", fontsize=6, style="italic", color=GRAY)
    arrow(ax, (xs[0], 23), (xs[1], 23), "14  POST /location", dy=0.7, fs=6.0)
    arrow(ax, (xs[1], 15), (xs[0], 15), "15  route state, guidance text", dy=0.7, fs=6.0)
    save(fig, "fig_sequence.png")


# ---------------------------------------------------------------------------------------
def slide_architecture():
    """The layered architecture picture from the team's presentation (slide 9)."""
    from pptx import Presentation
    deck = ROOT / "docs" / "SafeStride.pptx"
    if not deck.exists():
        return
    prs = Presentation(str(deck))
    for shape in prs.slides[8].shapes:
        if shape.shape_type == 13 and shape.image.size[1] > 1000:
            (OUT / "fig_architecture_layers.jpg").write_bytes(shape.image.blob)
            print("wrote fig_architecture_layers.jpg")
            return


# ---------------------------------------------------------------------------------------
def copy_existing():
    pa = ROOT / "docs" / "paper_assets"
    for src, dst in [("fig_scale.png", "fig_scale.png"), ("fig_hazard_map.png", "fig_hazard_map.png"),
                     ("fig_corridor.png", "fig_corridor.png"), ("fig_nav_states.png", "fig_nav_states.png"),
                     ("fig_live_frame.jpg", "fig_live_frame.jpg")]:
        if (pa / src).exists():
            shutil.copy(pa / src, OUT / dst)
    shot = Path(r"C:\Users\jhenk\.copilot\session-state\b3d46867-c283-4853-bb52-cb5b036cd0bd\files\console_shot.png")
    if shot.exists():
        shutil.copy(shot, OUT / "fig_console.png")


def notebook_figures():
    nb = json.loads((ROOT / "notebooks" / "train_hazards_results.ipynb").read_text(encoding="utf-8"))
    images = []
    for cell in nb["cells"]:
        for o in cell.get("outputs", []):
            if "data" in o and "image/png" in o["data"]:
                images.append(o["data"]["image/png"])
    for name, data in zip(("fig_train_curves.png", "fig_confusion.png"), images[-2:]):
        raw = data if isinstance(data, str) else "".join(data)
        (OUT / name).write_bytes(base64.b64decode(raw))
        print("wrote", name)


def failure_frame():
    import cv2
    import numpy as np
    sys.path.insert(0, str(ROOT / "src"))
    from blindnav.depth import DepthEstimator, colorize, object_distance
    from blindnav.detector import HAZARD_CLASS_CONF, Detector
    from blindnav.risk import MotionTracker
    from blindnav.spatial import SpatialAnalyzer, keep_confirmed_stairs
    from blindnav.visualize import draw_detections, draw_path

    detector = Detector("models/yolo11s.pt")
    hazards = Detector("models/hazards.pt", all_classes=True, class_conf=HAZARD_CLASS_CONF, id_offset=100_000)
    depth_model = DepthEstimator(input_size=392)
    cap = cv2.VideoCapture(str(ROOT / "data/videos/live_demo_source.mp4"))
    spatial, motion = SpatialAnalyzer(), MotionTracker()
    fps = cap.get(cv2.CAP_PROP_FPS) or 30
    best, n, depth = (-1, None), 0, None
    while n < int(10 * fps):
        ok, frame = cap.read()
        if not ok:
            break
        raw = frame.copy()
        dets = detector.detect(frame) + hazards.detect(frame)
        if depth is None or n % 2 == 0:
            depth = depth_model.estimate(frame)
        for d in dets:
            d.distance = object_distance(depth, d.box)
        status = spatial.analyze(depth, dets)
        dets = keep_confirmed_stairs(dets)
        motion.update(n / fps, dets, status, frame.shape[0])
        score = len([d for d in dets if d.class_name == "stairs"]) + (2 if status.blocked else 0)
        if n > 30 and score > best[0]:
            shown = raw.copy()
            draw_path(shown, status, dets)
            draw_detections(shown, dets)
            side = draw_detections(colorize(depth), dets)
            best = (score, np.hstack([shown, side]))
        n += 1
    image = best[1]
    scale = 1500 / image.shape[1]
    cv2.imwrite(str(OUT / "fig_failure_frame.jpg"), cv2.resize(image, None, fx=scale, fy=scale),
                [cv2.IMWRITE_JPEG_QUALITY, 90])
    print("wrote fig_failure_frame.jpg")


if __name__ == "__main__":
    parts = sys.argv[1:] or ["diagrams", "copy", "notebook", "failure"]
    if "diagrams" in parts:
        fig_hld(); fig_usecase_walk(); fig_usecase_route(); fig_dfd(); fig_sequence(); slide_architecture()
    if "copy" in parts:
        copy_existing()
    if "notebook" in parts:
        notebook_figures()
    if "failure" in parts:
        failure_frame()
