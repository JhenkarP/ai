"""Build the Safe-Stride Assist major-project report in the department's format.

    python docs/create_major_project_report.py [--template PATH] [--pages docs/report_assets/page_numbers.json]

The template (a report of another project in the same format) supplies the page setup, logos, headers,
footers, styles and bullet lists; its front matter is rewritten in place and its body is replaced by the
content of docs/report_content.py. Contents and list-of-figures page numbers are read from the pages file,
which scripts/report_pages.py measures by exporting the document to PDF with Microsoft Word.
"""

import argparse
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import report_content as C  # noqa: E402

ROOT = HERE.parent
ASSETS = HERE / "report_assets"
FONT = "Times New Roman"
J, CENTER, LEFT, RIGHT = (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER,
                          WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT)


# ---------------------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------------------
def style_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


TOKEN = re.compile(r"(\*\*.+?\*\*)")


def split_markup(text):
    parts = []
    for chunk in TOKEN.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            parts.append((chunk[2:-2], True))
        else:
            parts.append((chunk, False))
    return parts


def set_runs(paragraph, parts, base=0, size=None, align=None):
    """Replace the text of a template paragraph, keeping the formatting of one of its runs."""
    runs = paragraph.runs
    base_rpr = None
    if runs:
        src = runs[min(base, len(runs) - 1)]._r.rPr
        base_rpr = copy.deepcopy(src) if src is not None else None
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)
    for text, bold in parts:
        run = paragraph.add_run(text)
        if base_rpr is not None:
            run._r.insert(0, copy.deepcopy(base_rpr))
        run.font.bold = bold
        if size:
            run.font.size = Pt(size)
    if align is not None:
        paragraph.alignment = align


def set_text(paragraph, text, bold=None, **kw):
    if bold is None:
        bold = bool(paragraph.runs and paragraph.runs[0].bold)
    set_runs(paragraph, [(text, bold)], **kw)


def set_cell_text(cell, text):
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


# ---------------------------------------------------------------------------------------
class Report:
    def __init__(self, template, pages):
        self.doc = Document(str(template))
        self.pages = pages
        self.toc = []        # (kind, number, title, key)
        self.figs = []       # (number, title)
        self.first_chapter = True

    # ---------------------------------------------------------- front matter
    def front_matter(self, abstract):
        d = self.doc
        P = d.paragraphs
        team = C.TEAM
        names = [f"{n}  –  {u}" for n, u in team]
        for i, idx in enumerate((15, 16, 17, 18)):
            if i < len(names):
                set_text(P[idx], names[i], bold=True, align=CENTER)
            else:
                set_text(P[idx], "", bold=True, align=CENTER)
        set_text(P[20], "Under The Guidance Of", bold=False, align=CENTER)
        guide_name, guide_role = C.GUIDE
        set_text(P[22], guide_name, bold=True, align=CENTER, base=0, size=14)
        set_text(P[23], guide_role, bold=True, align=CENTER, base=1, size=14)
        set_text(P[24], "Department of CS&E,", bold=True, align=CENTER, size=14)
        set_text(P[25], "NIE Mysore", bold=True, align=CENTER, size=14)
        set_text(P[27], "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", bold=True)
        set_text(P[32], C.YEAR, bold=True)

        team_names = [n for n, _ in team]
        joined = team_names[0] if len(team_names) == 1 else ", ".join(team_names[:-1]) + " and " + team_names[-1]
        verb = "is a bonafide work carried out by "
        set_runs(P[47], [
            ("This is to certify that the project work entitled ", False),
            (f"“{C.TITLE}”", True), (f" {verb}", False), (joined, True),
            (" in fulfillment of the requirements for the major project work of the Bachelor of Engineering in "
             "Computer Science and Engineering, Seventh Semester, at The National Institute of Engineering "
             "(Autonomous under VTU), Mysuru, during the academic year ", False), (C.ACADEMIC_YEAR, True), (".", False)],
            base=0, size=12)

        # signature block: centre tab stops keep the three columns aligned whatever the name lengths
        cols = (0.95, 3.15, 5.3)
        rows = {
            52: (["Signature of Internal Guide", "Signature of HOD", "Signature of Principal"], True),
            55: (["_______________________"] * 3, False),
            57: ([guide_name, C.HOD[0], C.PRINCIPAL], False),
            58: ([guide_role, C.HOD[1], "Principal,"], False),
            59: (["Dept. of CS&E", "Dept. of CS&E", "NIE, Mysuru"], False),
            60: (["NIE, Mysuru", "NIE, Mysuru", ""], False),
        }
        for idx, (texts, bold) in rows.items():
            p = P[idx]
            p.paragraph_format.tab_stops.clear_all()
            for pos in cols:
                p.paragraph_format.tab_stops.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CENTER)
            p.alignment = LEFT
            set_runs(p, [("\t" + "\t".join(texts), bold)], base=0, size=12)

        for idx, text in zip((66, 67, 68, 69), abstract):
            set_runs(P[idx], split_markup(text), base=0, size=12, align=J)

        set_runs(P[82], [("First and foremost we would like to thank our beloved principal ", False),
                         (f"{C.PRINCIPAL} ", True), ("for being the patron and the beacon light for this project.", False)],
                 base=0, size=12)
        set_runs(P[84], [("We would like to express our sincere gratitude to our HOD ", False), (C.HOD[0], True),
                         (", Department of CSE, NIE for her relentless support and encouragement. It gives us immense "
                          "pleasure to thank our guide ", False), (guide_name, True),
                         (", Assistant Professor, Department of CSE, NIE for the valuable suggestions and guidance "
                          "during the process of the project and for having permitted us to pursue work on the subject.", False)],
                 base=0, size=12)
        for i, idx in enumerate((89, 90, 91, 92)):
            text = f"{team[i][0]} ({team[i][1]})" if i < len(team) else ""
            set_runs(P[idx], [(text, True)], base=0, size=12, align=RIGHT)

        for i, sec in enumerate(d.sections):
            for para in sec.header.paragraphs:
                if para.runs and para.runs[0].text.strip():
                    para.runs[0].text = C.HEADER
                    if i < 2:
                        para.runs[0].font.size = Pt(12)  # the longer title must still fit two lines
        # the third section has its own wording of the header
        for sec in d.sections:
            if not sec.header.is_linked_to_previous:
                for para in sec.header.paragraphs:
                    if para.runs and para.runs[0].text.strip():
                        para.runs[0].text = C.HEADER
        # each front-matter item starts on its own page, as in the template
        P[64].paragraph_format.page_break_before = True
        P[78].paragraph_format.page_break_before = True
        # a trailing empty paragraph after the examiners line would otherwise spill onto a blank page
        for idx in (63,):
            el = P[idx]._p
            if not P[idx].text.strip() and el.find(qn("w:pPr")) is not None \
                    and el.find(qn("w:pPr")).find(qn("w:sectPr")) is None and not el.findall(".//" + qn("w:drawing")):
                el.getparent().remove(el)

    def cut_old_body(self, anchor):
        d = self.doc
        assert anchor.find(qn("w:pPr")) is not None and anchor.find(qn("w:pPr")).find(qn("w:sectPr")) is not None, \
            "template layout changed: expected the section break at paragraph 131"
        sib = anchor.getnext()
        body = d.element.body
        while sib is not None and sib.tag != qn("w:sectPr"):
            nxt = sib.getnext()
            body.remove(sib)
            sib = nxt

    # ---------------------------------------------------------- paragraph builders
    def par(self, parts, size=12, align=J, before=0, after=8, line=1.5, left=None, first=None,
            keep_next=False, page_break=False, italic=False, bold=False, keep_lines=False):
        p = self.doc.add_paragraph()
        f = p.paragraph_format
        p.alignment = align
        f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
        if left is not None:
            f.left_indent = Inches(left)
        if first is not None:
            f.first_line_indent = Inches(first)
        f.keep_with_next = keep_next
        f.keep_together = keep_lines
        f.page_break_before = page_break
        if isinstance(parts, str):
            parts = split_markup(parts)
        for text, b in parts:
            style_run(p.add_run(text), size, bold or b, italic)
        return p

    def chapter(self, n, title):
        self.par([(f"Chapter {n}", True)], size=16, align=LEFT, before=0, after=4, line=1.0,
                 page_break=not self.first_chapter, keep_next=True)
        self.first_chapter = False
        p = self.doc.add_paragraph(style="Heading 1")
        p.alignment = CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(14)
        p.paragraph_format.keep_with_next = True
        style_run(p.add_run(title), 14, True)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
        self.toc.append(("chapter", str(n), title.title().replace(" And ", " and "), f"ch{n}"))

    def h2(self, num, title):
        self.par([(f"{num}  {title}", True)], size=14, align=LEFT, before=16, after=8, line=1.0, keep_next=True)
        self.toc.append(("section", num, title, f"s{num}"))

    def h3(self, num, title):
        self.par([(f"{num}  {title}", True)], size=13, align=LEFT, before=12, after=6, line=1.0, keep_next=True)

    def plain_h(self, text):
        self.par([(text, True)], size=13, align=LEFT, before=12, after=6, line=1.0, keep_next=True)

    def body(self, text):
        self.par(text)

    def bullets(self, items):
        for lead, rest in items:
            p = self.doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            num = ppr.get_or_add_numPr()
            num.get_or_add_ilvl().val = 0
            num.get_or_add_numId().val = 7
            f = p.paragraph_format
            p.alignment = J
            f.space_before, f.space_after, f.line_spacing = Pt(0), Pt(4), 1.4
            f.left_indent, f.first_line_indent = Inches(0.5), Inches(-0.25)
            if lead:
                style_run(p.add_run(lead), 12, True)
            for text, b in split_markup(rest):
                style_run(p.add_run(text), 12, b)
        self.par("", after=2, line=1.0)

    def numbered(self, items):
        for i, (lead, text) in enumerate(items, 1):
            self.par([(f"{i}. {lead}: ", True), (text, False)], after=8, line=1.4)

    def centered(self, text):
        self.par([(text, False)], align=CENTER, before=8, after=6, line=1.0, keep_next=True)

    def lit(self, n, title, text):
        self.par([(f"{n}. {title}", True)], size=12.5, align=LEFT, before=10, after=4, line=1.15, keep_next=True)
        self.par(text, after=8, line=1.4)

    def fig(self, file, caption, width):
        path = ASSETS / file
        p = self.doc.add_paragraph()
        p.alignment = CENTER
        f = p.paragraph_format
        f.space_before, f.space_after, f.line_spacing = Pt(8), Pt(2), 1.0
        f.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        self.par([(caption, False)], size=11.5, align=CENTER, before=2, after=12, line=1.0, italic=True)
        m = re.match(r"(Fig \d+\.\d+): (.*)", caption)
        if m:
            self.figs.append((m.group(1), m.group(2)))

    def code(self, title, lines):
        self.par([(title, True)], size=13, align=LEFT, before=14, after=4, line=1.0, keep_next=True)
        self.par([("Pseudocode:", True)], size=12, align=LEFT, before=0, after=4, line=1.0, keep_next=True)
        for level, text in lines:
            p = self.doc.add_paragraph()
            f = p.paragraph_format
            p.alignment = LEFT
            f.space_before, f.space_after, f.line_spacing = Pt(0), Pt(1), 1.1
            f.left_indent = Inches(0.25 + 0.3 * level)
            f.first_line_indent = Inches(-0.0)
            style_run(p.add_run(text), 11, text in ("BEGIN", "END"))

    def table(self, caption, header, rows, widths, align):
        self.par([(caption, False)], size=11.5, align=CENTER, before=10, after=4, line=1.0, keep_next=True, bold=False)
        total = sum(widths)
        scale = min(6.2, total) / total
        widths = [w * scale for w in widths]
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
                el.set(qn(f"w:{k}"), v)
            borders.append(el)
        look = tblPr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(borders)
        else:
            tblPr.append(borders)
        for ri, row in enumerate([header] + rows):
            tr = tbl.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for ci, value in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(widths[ci])
                para = cell.paragraphs[0]
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                text_col = align == "left" and ci > 0 and ri > 0
                para.alignment = LEFT if (text_col or (ri == 0 and ci == 0 and align == "left")) else CENTER
                if align == "left" and ri > 0 and ci == 0 and not str(value).isdigit():
                    para.alignment = LEFT
                style_run(para.add_run(str(value)), 10.5, ri == 0)
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D9D9D9")
                    tcPr.append(shd)
        self.par("", after=6, line=1.0)

    def refs(self, items):
        self.par([("REFERENCES", True)], size=14, align=LEFT, before=0, after=10, line=1.0, page_break=True, keep_next=True)
        self.toc.append(("chapter", "8", "References", "refs"))
        for n, text in enumerate(items, 1):
            p = self.par(f"[{n}] {text}", size=12, after=6, line=1.15, left=0.4, first=-0.4)

    # ---------------------------------------------------------- tables of contents
    def fill_toc(self):
        d = self.doc
        toc_tbl, fig_tbl = d.tables[0], d.tables[1]
        head, chap, sect = (copy.deepcopy(toc_tbl.rows[i]._tr) for i in (0, 1, 2))
        for r in list(toc_tbl.rows):
            toc_tbl._tbl.remove(r._tr)
        toc_tbl._tbl.append(head)
        for kind, num, title, key in self.toc:
            tr = copy.deepcopy(chap if kind == "chapter" else sect)
            toc_tbl._tbl.append(tr)
            row = toc_tbl.rows[-1]
            label = f"{num}.\t{title}"
            set_cell_text(row.cells[0], label)
            set_cell_text(row.cells[1], str(self.pages.get(key, "")))

        fhead, frow = (copy.deepcopy(fig_tbl.rows[i]._tr) for i in (0, 1))
        for r in list(fig_tbl.rows):
            fig_tbl._tbl.remove(r._tr)
        fig_tbl._tbl.append(fhead)
        for number, title in self.figs:
            fig_tbl._tbl.append(copy.deepcopy(frow))
            row = fig_tbl.rows[-1]
            set_cell_text(row.cells[0], number)
            set_cell_text(row.cells[1], C.LOF_SHORT.get(number, title))
            set_cell_text(row.cells[2], str(self.pages.get(number, "")))

    # ---------------------------------------------------------- driver
    def render(self, blocks):
        abstract = next(b[1] for b in blocks if b[0] == "abstract")
        anchor = self.doc.paragraphs[131]._p  # section break that closes the list of figures
        self.front_matter(abstract)
        self.cut_old_body(anchor)
        for b in blocks:
            kind = b[0]
            if kind == "abstract":
                continue
            getattr(self, {"p": "body"}.get(kind, kind))(*b[1:])
        self.fill_toc()

    def save(self, path):
        p = Path(path)
        for attempt in range(10):
            target = p if attempt == 0 else p.with_name(f"{p.stem}_v{attempt + 1}{p.suffix}")
            try:
                self.doc.core_properties.title = C.TITLE
                self.doc.core_properties.author = ", ".join(n for n, _ in C.TEAM)
                self.doc.save(str(target))
                print("Wrote", target)
                return target
            except PermissionError:
                print(f"{target.name} is open in another program; trying a new name")
        raise PermissionError(p)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", default=str(Path.home() / "Downloads" / "Copy of Copy of major_rep.docx"))
    ap.add_argument("--pages", default=str(ASSETS / "page_numbers.json"))
    ap.add_argument("--out", default=str(HERE / "SafeStride_Assist_Major_Project_Report.docx"))
    args = ap.parse_args()
    pages = json.loads(Path(args.pages).read_text(encoding="utf-8")) if Path(args.pages).exists() else {}
    results = json.loads((HERE / "paper_assets" / "results.json").read_text(encoding="utf-8"))
    report = Report(args.template, pages)
    report.render(C.build(results))
    report.save(args.out)


if __name__ == "__main__":
    main()
