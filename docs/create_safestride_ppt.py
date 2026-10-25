from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("SafeStride_Assist_Presentation.pptx")

NAVY = RGBColor(13, 21, 28)
PANEL = RGBColor(24, 35, 42)
MINT = RGBColor(184, 231, 105)
CYAN = RGBColor(101, 212, 196)
WHITE = RGBColor(239, 243, 245)
MUTED = RGBColor(160, 174, 181)
RED = RGBColor(255, 114, 114)
AMBER = RGBColor(242, 187, 98)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=PANEL, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def bullets(slide, items, x, y, w, h, size=18, color=WHITE, accent=MINT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = "- " + item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
    return shape


def base(title, kicker, number):
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    box(slide, 0, 0, 0.18, 7.5, fill=MINT)
    text(slide, kicker.upper(), 0.65, 0.38, 5.0, 0.25, size=10, color=CYAN, bold=True, font="Aptos Mono")
    text(slide, title, 0.65, 0.73, 11.8, 0.55, size=28, color=WHITE, bold=True)
    text(slide, f"SAFESTRIDE ASSIST  /  {number:02d}", 10.8, 7.08, 2.0, 0.2, size=9, color=MUTED, font="Aptos Mono", align=PP_ALIGN.RIGHT)
    return slide


def pill(slide, label, x, y, w, color=MINT):
    box(slide, x, y, w, 0.34, fill=color, line=color, radius=True)
    text(slide, label, x, y + 0.04, w, 0.22, size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, color=MINT):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


# 1 title
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 0.22, 7.5, fill=MINT)
text(slide, "SAFESTRIDE", 0.8, 0.7, 4.5, 0.4, size=16, color=MINT, bold=True, font="Aptos Mono")
text(slide, "Assistive vision\nand navigation", 0.78, 1.55, 8.2, 1.55, size=42, color=WHITE, bold=True)
text(slide, "A real-time research prototype for distance-aware obstacle guidance, GPS routing, and voice feedback.", 0.82, 3.45, 7.6, 0.75, size=20, color=MUTED)
box(slide, 9.0, 1.15, 3.35, 4.65, fill=PANEL, line=CYAN, radius=True)
text(slide, "LOCAL\nPERCEPTION", 9.35, 1.6, 2.6, 0.8, size=20, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
arrow(slide, 10.65, 2.62, 10.65, 3.18, MINT)
text(slide, "DEPTH +\nCORRIDOR", 9.35, 3.35, 2.6, 0.8, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
arrow(slide, 10.65, 4.35, 10.65, 4.9, MINT)
text(slide, "VOICE\nGUIDANCE", 9.35, 5.05, 2.6, 0.8, size=20, color=MINT, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Major Project Synopsis + Presentation", 0.82, 6.55, 6.5, 0.3, size=13, color=MUTED)

# 2 problem
slide = base("The problem", "Why this matters", 2)
box(slide, 0.7, 1.65, 5.6, 4.7, fill=PANEL, radius=True)
text(slide, "Object detection alone is not enough", 1.05, 2.0, 4.8, 0.6, size=24, color=MINT, bold=True)
bullets(slide, [
    "A box does not say whether an object is in the walking path.",
    "Users need distance, direction, urgency, and an actionable next step.",
    "Repeated low-priority alerts can hide a genuinely urgent obstruction.",
    "Route guidance and local safety perception must work together."
], 1.05, 2.9, 4.75, 2.6, size=18)
box(slide, 7.0, 1.65, 5.55, 4.7, fill=RGBColor(30, 43, 39), line=CYAN, radius=True)
text(slide, "SafeStride's answer", 7.4, 2.0, 4.8, 0.45, size=24, color=CYAN, bold=True)
text(slide, "Perceive  ->  reason  ->  prioritise  ->  speak", 7.4, 2.85, 4.7, 0.45, size=20, color=WHITE, bold=True)
text(slide, "The prototype translates visual evidence into concise guidance such as:", 7.4, 3.65, 4.5, 0.55, size=16, color=MUTED)
box(slide, 7.4, 4.55, 4.65, 0.85, fill=NAVY, line=MINT, radius=True)
text(slide, "Path blocked, 2 meters.\nKeep left.", 7.65, 4.73, 4.15, 0.5, size=20, color=MINT, bold=True)

# 3 objectives
slide = base("Project objectives", "What we built", 3)
objectives = [
    ("01", "See", "Detect navigation-relevant objects and unknown depth obstacles."),
    ("02", "Measure", "Estimate distance, clock direction, approach risk, and path position."),
    ("03", "Guide", "Give complete voice instructions with left/right corrections."),
    ("04", "Navigate", "Combine GPS walking routes with local camera safety."),
    ("05", "Respect", "Keep processing local and expose clear prototype limitations."),
]
for i, (num, head, body) in enumerate(objectives):
    y = 1.55 + i * 1.02
    text(slide, num, 0.85, y, 0.55, 0.35, size=14, color=MINT, bold=True, font="Aptos Mono")
    text(slide, head, 1.65, y - 0.03, 1.7, 0.35, size=21, color=WHITE, bold=True)
    text(slide, body, 3.35, y, 8.9, 0.38, size=17, color=MUTED)
    box(slide, 0.82, y + 0.62, 11.5, 0.02, fill=RGBColor(47, 63, 68))

# 4 architecture
slide = base("System architecture", "One pipeline, two guidance scales", 4)
labels = [
    ("IP camera / video", 0.75, 2.0, 2.0, CYAN),
    ("YOLO + ByteTrack", 3.15, 2.0, 2.0, MINT),
    ("Depth Anything V2", 5.55, 2.0, 2.0, AMBER),
    ("3D corridor\nreasoning", 7.95, 2.0, 2.0, RED),
    ("Alert + route\nstate", 10.35, 2.0, 2.0, CYAN),
]
for label, x, y, w, color in labels:
    box(slide, x, y, w, 1.0, fill=PANEL, line=color, radius=True)
    text(slide, label, x + 0.12, y + 0.25, w - 0.24, 0.5, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
for i in range(len(labels) - 1):
    arrow(slide, labels[i][1] + labels[i][3], 2.5, labels[i + 1][1], 2.5, MINT)
box(slide, 1.0, 4.1, 5.4, 1.35, fill=RGBColor(30, 43, 39), line=MINT, radius=True)
text(slide, "LOCAL SAFETY", 1.35, 4.38, 1.8, 0.25, size=12, color=MINT, bold=True, font="Aptos Mono")
text(slide, "Immediate obstruction, distance, side correction", 1.35, 4.78, 4.5, 0.3, size=17, color=WHITE)
box(slide, 6.9, 4.1, 5.4, 1.35, fill=RGBColor(30, 39, 48), line=CYAN, radius=True)
text(slide, "MACRO ROUTE", 7.25, 4.38, 1.8, 0.25, size=12, color=CYAN, bold=True, font="Aptos Mono")
text(slide, "GPS position, turns, deviation, arrival", 7.25, 4.78, 4.5, 0.3, size=17, color=WHITE)
text(slide, "Urgent camera safety always wins over routine route speech.", 2.6, 6.2, 8.2, 0.35, size=16, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

# 5 perception
slide = base("Perception and geometry", "How the system understands a frame", 5)
steps = [
    ("1", "Detect", "YOLO finds relevant objects."),
    ("2", "Track", "ByteTrack keeps identity over time."),
    ("3", "Estimate", "Depth Anything V2 gives relative depth."),
    ("4", "Calibrate", "Ground plane converts depth to approximate metres."),
    ("5", "Decide", "A 1 m wide corridor is checked up to 3 m ahead."),
]
for i, (num, head, body) in enumerate(steps):
    x = 0.75 + i * 2.48
    box(slide, x, 1.8, 2.05, 2.9, fill=PANEL, line=[CYAN, MINT, AMBER, RED, MINT][i], radius=True)
    text(slide, num, x + 0.22, 2.1, 0.4, 0.35, size=16, color=MINT, bold=True, font="Aptos Mono")
    text(slide, head, x + 0.22, 2.7, 1.6, 0.4, size=21, color=WHITE, bold=True)
    text(slide, body, x + 0.22, 3.45, 1.6, 0.85, size=15, color=MUTED)
    if i < len(steps) - 1:
        arrow(slide, x + 2.08, 3.25, x + 2.4, 3.25, MINT)
box(slide, 2.0, 5.45, 9.3, 0.8, fill=RGBColor(47, 35, 35), line=RED, radius=True)
text(slide, "PATH BLOCKED  ->  find nearest point  ->  choose clear side  ->  speak distance + action", 2.25, 5.68, 8.8, 0.3, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 6 alerts
slide = base("Safety-first alerting", "Turning evidence into action", 6)
box(slide, 0.8, 1.65, 4.0, 4.7, fill=PANEL, radius=True)
text(slide, "Priority order", 1.2, 2.0, 3.2, 0.4, size=24, color=MINT, bold=True)
for i, (label, color) in enumerate([
    ("Urgent obstruction", RED), ("Local correction", AMBER), ("Route instruction", CYAN), ("Background status", MUTED)
]):
    y = 2.85 + i * 0.72
    box(slide, 1.2, y, 3.0, 0.46, fill=color, line=color, radius=True)
    text(slide, label, 1.2, y + 0.1, 3.0, 0.25, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
box(slide, 5.35, 1.65, 7.15, 4.7, fill=RGBColor(30, 43, 39), line=CYAN, radius=True)
text(slide, "Example interaction", 5.8, 2.0, 3.4, 0.4, size=24, color=CYAN, bold=True)
text(slide, "Camera sees an obstacle in the corridor.", 5.8, 2.9, 5.8, 0.35, size=18, color=WHITE)
arrow(slide, 8.7, 3.45, 8.7, 3.9, MINT)
text(slide, "Depth estimates 2 m;\nspatial analysis finds left escape.", 5.8, 4.0, 5.8, 0.65, size=18, color=WHITE)
arrow(slide, 8.7, 4.85, 8.7, 5.25, MINT)
box(slide, 5.8, 5.35, 5.8, 0.65, fill=NAVY, line=MINT, radius=True)
text(slide, "Stop. Path blocked, 2 meters. Keep left.", 6.0, 5.55, 5.4, 0.25, size=17, color=MINT, bold=True, align=PP_ALIGN.CENTER)

# 7 navigation
slide = base("Navigation layer", "From location to route-aware guidance", 7)
route_steps = ["Browser GPS", "Geocode\ndestination", "Walking route", "Turn state", "Voice prompt"]
for i, label in enumerate(route_steps):
    x = 0.85 + i * 2.45
    box(slide, x, 2.15, 1.95, 1.05, fill=PANEL, line=CYAN if i != 3 else AMBER, radius=True)
    text(slide, label, x + 0.12, 2.43, 1.7, 0.45, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if i < len(route_steps) - 1:
        arrow(slide, x + 1.98, 2.68, x + 2.35, 2.68, CYAN)
box(slide, 1.25, 4.25, 10.7, 1.15, fill=RGBColor(30, 39, 48), line=CYAN, radius=True)
text(slide, "NavigationManager", 1.7, 4.55, 2.5, 0.3, size=20, color=CYAN, bold=True)
text(slide, "route geometry  /  turn instructions  /  off-route detection  /  arrival  /  repeated prompts", 4.0, 4.55, 7.3, 0.3, size=16, color=WHITE)
text(slide, "Route speech is paused briefly when a camera obstruction alert is active.", 2.05, 6.05, 9.2, 0.3, size=16, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

# 8 interface
slide = base("Live console", "A practical demonstration surface", 8)
box(slide, 0.75, 1.55, 7.4, 4.95, fill=PANEL, line=CYAN, radius=True)
box(slide, 1.05, 1.95, 6.8, 3.25, fill=RGBColor(9, 14, 18), line=RGBColor(47, 63, 68), radius=True)
# simple camera mockup
for x in [1.35, 2.1, 2.85, 3.6, 4.35, 5.1, 5.85, 6.6, 7.25]:
    box(slide, x, 2.3, 0.08, 2.55, fill=RGBColor(80, 138, 106))
box(slide, 1.35, 4.55, 6.2, 0.38, fill=RGBColor(46, 20, 22), line=RED)
text(slide, "PATH BLOCKED: obstacle 1.9 m  KEEP LEFT", 1.6, 4.63, 5.7, 0.2, size=12, color=RED, bold=True, font="Aptos Mono", align=PP_ALIGN.CENTER)
for i, (label, value) in enumerate([("Objects", "2"), ("Inference", "27 fps"), ("Depth", "ACTIVE")]):
    x = 8.65 + (i % 1) * 0.0
    y = 1.7 + i * 1.15
    box(slide, x, y, 3.7, 0.85, fill=PANEL, line=RGBColor(47, 63, 68), radius=True)
    text(slide, label, x + 0.22, y + 0.14, 1.1, 0.2, size=11, color=MUTED, font="Aptos Mono")
    text(slide, value, x + 1.45, y + 0.13, 1.9, 0.28, size=19, color=MINT, bold=True)
box(slide, 8.65, 5.25, 3.7, 0.85, fill=RGBColor(30, 43, 39), line=MINT, radius=True)
text(slide, "Route: active\nVoice: enabled", 8.9, 5.42, 3.2, 0.45, size=16, color=WHITE, bold=True)

# 9 datasets/extensions
slide = base("Model extensions", "Doors and stairs", 9)
box(slide, 0.8, 1.6, 5.55, 4.75, fill=PANEL, line=CYAN, radius=True)
text(slide, "Door state", 1.2, 1.98, 3.0, 0.4, size=25, color=CYAN, bold=True)
bullets(slide, [
    "door_open -> announce open door",
    "door_closed -> obstruction warning",
    "door_semi -> partial obstruction warning",
    "Distance gate: only 0–3 metres"
], 1.2, 2.85, 4.55, 2.25, size=17)
text(slide, "Dataset starting point: door-open-close\n687 object-detection images", 1.2, 5.55, 4.6, 0.45, size=13, color=MUTED)
box(slide, 6.9, 1.6, 5.55, 4.75, fill=PANEL, line=AMBER, radius=True)
text(slide, "Stair direction", 7.3, 1.98, 3.5, 0.4, size=25, color=AMBER, bold=True)
bullets(slide, [
    "stair_up -> stairs going up",
    "stair_down -> stairs going down",
    "Depth supplies distance and path relevance",
    "Evaluate on the project's camera viewpoint"
], 7.3, 2.85, 4.55, 2.25, size=17)
text(slide, "Dataset starting point: stair-up-or-down\n923 object-detection images", 7.3, 5.55, 4.6, 0.45, size=13, color=MUTED)

# 10 evaluation
slide = base("Evaluation plan", "Evidence before confidence", 10)
metrics = [("Precision / recall", "Can it detect the right things?"), ("mAP50", "How strong are the boxes?"), ("Distance error", "How close is the metre estimate?"), ("False alerts", "Does it remain quiet when clear?"), ("Latency / FPS", "Can it respond in real time?")]
for i, (head, body) in enumerate(metrics):
    y = 1.55 + i * 0.92
    box(slide, 0.9, y, 3.05, 0.55, fill=MINT, line=MINT, radius=True)
    text(slide, head, 0.9, y + 0.14, 3.05, 0.22, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, 4.35, y + 0.13, 6.9, 0.25, size=17, color=WHITE)
text(slide, "Test groups", 8.9, 1.25, 2.2, 0.3, size=15, color=CYAN, bold=True, font="Aptos Mono")
for i, label in enumerate(["people", "vehicles", "stairs", "bollards", "crowds", "low light"]):
    x = 8.9 + (i % 2) * 1.65
    y = 1.8 + (i // 2) * 0.75
    pill(slide, label, x, y, 1.4, CYAN if i % 2 else MINT)
box(slide, 8.9, 4.55, 3.3, 1.05, fill=RGBColor(47, 35, 35), line=RED, radius=True)
text(slide, "Prototype only\nnot a certified mobility aid", 9.15, 4.78, 2.8, 0.5, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 11 status
slide = base("Current status and next steps", "Where the project goes next", 11)
box(slide, 0.8, 1.6, 5.6, 4.75, fill=RGBColor(30, 43, 39), line=MINT, radius=True)
text(slide, "Working now", 1.2, 1.98, 3.0, 0.4, size=25, color=MINT, bold=True)
bullets(slide, [
    "Live IP-camera browser console",
    "YOLO + depth + corridor analysis",
    "Distance-aware voice alerts",
    "GPS route and navigation state",
    "Safety priority over route speech"
], 1.2, 2.85, 4.7, 2.55, size=18)
box(slide, 6.95, 1.6, 5.6, 4.75, fill=PANEL, line=CYAN, radius=True)
text(slide, "Next", 7.35, 1.98, 3.0, 0.4, size=25, color=CYAN, bold=True)
bullets(slide, [
    "Train door open / closed / semi-closed",
    "Train stair-up / stair-down",
    "Collect project-specific camera data",
    "Benchmark false alerts and distance error",
    "Improve offline and mobile support"
], 7.35, 2.85, 4.7, 2.55, size=18)

# 12 close
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
box(slide, 0, 0, 0.22, 7.5, fill=MINT)
text(slide, "SAFESTRIDE ASSIST", 0.85, 1.0, 5.0, 0.35, size=15, color=MINT, bold=True, font="Aptos Mono")
text(slide, "Perception becomes\nguidance when it is\nplaced in context.", 0.82, 1.85, 7.4, 2.05, size=38, color=WHITE, bold=True)
text(slide, "A local-processing prototype for safer, more understandable navigation assistance.", 0.85, 4.55, 6.5, 0.45, size=19, color=MUTED)
box(slide, 9.05, 1.7, 3.0, 3.0, fill=PANEL, line=MINT, radius=True)
text(slide, "SEE", 9.45, 2.1, 2.2, 0.35, size=22, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "REASON", 9.45, 2.95, 2.2, 0.35, size=22, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
text(slide, "SPEAK", 9.45, 3.8, 2.2, 0.35, size=22, color=MINT, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Thank you", 0.85, 6.6, 2.0, 0.3, size=15, color=MUTED)

prs.save(OUT)
print(OUT)
