from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("SafeStride_B15_Presentation.pptx")
ICON = Path(__file__).parents[1] / "mobile" / "assets" / "icon.png"

NAVY = RGBColor(13, 21, 28)
PANEL = RGBColor(24, 35, 42)
MINT = RGBColor(184, 231, 105)
CYAN = RGBColor(101, 212, 196)
WHITE = RGBColor(239, 243, 245)
MUTED = RGBColor(160, 174, 181)
RED = RGBColor(255, 114, 114)
AMBER = RGBColor(242, 187, 98)
LINE = RGBColor(47, 63, 68)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill=PANEL, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_bullets(slide, items, x, y, w, h, size=17, color=WHITE, gap=7):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = "- " + item
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(gap)
    return shape


def base(title, kicker, number):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.18, 7.5, fill=MINT)
    add_text(slide, kicker.upper(), 0.65, 0.32, 6, 0.24, size=10, color=CYAN, bold=True, font="Aptos Mono")
    add_text(slide, title, 0.65, 0.68, 11.7, 0.52, size=27, color=WHITE, bold=True)
    add_text(slide, f"SAFESTRIDE ASSIST  /  {number:02d}", 10.6, 7.08, 2.2, 0.2, size=9, color=MUTED, font="Aptos Mono", align=PP_ALIGN.RIGHT)
    return slide


def arrow(slide, x1, y1, x2, y2, color=MINT):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def pill(slide, label, x, y, w, color=MINT):
    rect(slide, x, y, w, 0.34, fill=color, line=color, rounded=True)
    add_text(slide, label, x, y + 0.04, w, 0.22, size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)


def code_block(slide, code, x, y, w, h, title):
    rect(slide, x, y, w, h, fill=RGBColor(9, 14, 18), line=LINE, rounded=True)
    add_text(slide, title.upper(), x + 0.18, y + 0.16, w - 0.36, 0.2, size=10, color=CYAN, bold=True, font="Aptos Mono")
    add_text(slide, code, x + 0.18, y + 0.52, w - 0.36, h - 0.65, size=12, color=WHITE, font="Aptos Mono")


# 1. Starting slide
slide = prs.slides.add_slide(BLANK)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
rect(slide, 0, 0, 0.22, 7.5, fill=MINT)
if ICON.exists():
    slide.shapes.add_picture(str(ICON), Inches(10.3), Inches(0.55), width=Inches(1.7), height=Inches(1.7))
add_text(slide, "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING", 0.8, 0.55, 8.8, 0.3, size=13, color=WHITE, bold=True)
add_text(slide, "THE NATIONAL INSTITUTE OF ENGINEERING", 0.8, 0.9, 8.8, 0.3, size=13, color=WHITE, bold=True)
add_text(slide, "(An Autonomous Institution)", 0.8, 1.25, 8.8, 0.25, size=11, color=MUTED)
add_text(slide, "PROJECT SYNOPSIS", 0.8, 1.85, 5.2, 0.35, size=15, color=MINT, bold=True, font="Aptos Mono")
add_text(slide, "SafeStride Assist", 0.8, 2.35, 8.2, 0.7, size=38, color=WHITE, bold=True)
add_text(slide, "An AI-Based Safe Walking Assistant Using Monocular Depth Estimation", 0.82, 3.18, 7.5, 0.65, size=22, color=CYAN, bold=True)
add_text(slide, "Batch B15  |  Jhenkar P  |  Major Project", 0.82, 4.25, 6.6, 0.3, size=16, color=MUTED)
rect(slide, 8.55, 2.35, 3.5, 2.9, fill=PANEL, line=CYAN, rounded=True)
add_text(slide, "SEE", 9.0, 2.78, 2.6, 0.35, size=24, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
arrow(slide, 10.3, 3.2, 10.3, 3.55)
add_text(slide, "REASON", 9.0, 3.7, 2.6, 0.35, size=24, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
arrow(slide, 10.3, 4.12, 10.3, 4.47)
add_text(slide, "SPEAK", 9.0, 4.62, 2.6, 0.35, size=24, color=MINT, bold=True, align=PP_ALIGN.CENTER)

# 2. Abstract
slide = base("Abstract", "Project overview", 2)
rect(slide, 0.8, 1.55, 7.1, 4.9, fill=PANEL, line=CYAN, rounded=True)
add_text(slide, "SafeStride Assist is a research and teaching prototype that combines computer vision, monocular depth estimation, spatial reasoning, and speech feedback to support safer walking for people with visual impairment.", 1.15, 1.95, 6.4, 1.05, size=21, color=WHITE, bold=True)
add_text(slide, "A camera stream is processed locally using YOLO-based detection, ByteTrack tracking, and Depth Anything V2. A ground-plane model converts relative depth into approximate metres and evaluates a three-metre forward walking corridor.", 1.15, 3.25, 6.35, 1.0, size=17, color=MUTED)
add_text(slide, "The prototype reports named obstacles and unknown depth-based blockers through a browser console, directional voice alerts, and approach-risk estimation. It is not a certified medical, mobility, or emergency device.", 1.15, 4.55, 6.35, 1.05, size=17, color=MUTED)
rect(slide, 8.45, 1.55, 3.9, 4.9, fill=RGBColor(30, 43, 39), line=MINT, rounded=True)
for i, (label, value, color) in enumerate([("Input", "Camera / video", CYAN), ("Reasoning", "Depth + corridor", AMBER), ("Output", "Voice guidance", MINT)]):
    y = 2.05 + i * 1.35
    add_text(slide, label.upper(), 8.85, y, 3.1, 0.22, size=10, color=color, bold=True, font="Aptos Mono", align=PP_ALIGN.CENTER)
    add_text(slide, value, 8.85, y + 0.38, 3.1, 0.35, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 3. Index
slide = base("Index", "Presentation roadmap", 3)
items = [
    ("01", "Introduction and motivation"),
    ("02", "Problem statement"),
    ("03", "Literature review"),
    ("04", "System architecture"),
    ("05", "System design and workflow"),
    ("06", "Algorithm code snippets"),
    ("07", "Implementation and safety"),
    ("08", "Novelty and contribution"),
    ("09", "Conclusion"),
    ("10", "References"),
]
for i, (number, label) in enumerate(items):
    x = 0.9 + (i // 5) * 6.1
    y = 1.45 + (i % 5) * 0.9
    add_text(slide, number, x, y, 0.65, 0.3, size=14, color=MINT, bold=True, font="Aptos Mono")
    add_text(slide, label, x + 0.9, y - 0.02, 4.7, 0.35, size=19, color=WHITE, bold=True)
    rect(slide, x, y + 0.48, 5.25, 0.02, fill=LINE)

# 4. Introduction
slide = base("Introduction", "Why safe walking assistance matters", 4)
add_text(slide, "A camera can provide an additional perception channel, but it must turn visual evidence into a useful walking decision.", 0.85, 1.45, 11.6, 0.55, size=25, color=MINT, bold=True)
intro = [
    "People with visual impairment may combine a white cane, guide dog, personal knowledge, and human assistance in unfamiliar environments.",
    "Traditional aids remain essential, while an experimental camera assistant can add information about nearby blockers and approximate distance.",
    "A visible object is not automatically a relevant hazard: path position, distance, height, persistence, and approach risk matter.",
    "SafeStride Assist is designed as an additional perception and feedback layer, not a replacement for conventional mobility support.",
]
for i, item in enumerate(intro):
    y = 2.45 + i * 0.78
    rect(slide, 0.95, y, 0.4, 0.4, fill=[CYAN, MINT, AMBER, RED][i], rounded=True)
    add_text(slide, str(i + 1), 0.95, y + 0.07, 0.4, 0.2, size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, item, 1.65, y - 0.03, 10.4, 0.52, size=18, color=WHITE)

# 5. Problem statement
slide = base("Problem statement", "From detection to decision", 5)
rect(slide, 0.8, 1.55, 5.65, 4.8, fill=PANEL, line=RED, rounded=True)
add_text(slide, "A bounding box is not a walking instruction", 1.2, 1.98, 4.8, 0.7, size=25, color=RED, bold=True)
add_bullets(slide, [
    "The detector may not know whether an object blocks the forward path.",
    "A label does not provide reliable distance, direction, or urgency.",
    "Unknown structures may not have a trained object class.",
    "Repeated low-priority speech can hide urgent warnings.",
], 1.2, 3.0, 4.6, 2.2, size=17)
rect(slide, 6.95, 1.55, 5.5, 4.8, fill=RGBColor(30, 43, 39), line=MINT, rounded=True)
add_text(slide, "Project objective", 7.35, 1.98, 3.8, 0.4, size=25, color=MINT, bold=True)
add_text(slide, "Design a real-time prototype that can:", 7.35, 2.7, 4.3, 0.35, size=18, color=WHITE)
add_bullets(slide, [
    "detect named obstacles and unknown depth blockers;",
    "estimate approximate distance and clock direction;",
    "decide whether a three-metre corridor is blocked;",
    "produce concise, prioritised, actionable voice guidance.",
], 7.35, 3.25, 4.2, 2.2, size=17)

# 6. Literature review
slide = base("Literature review", "Research foundation", 6)
research = [
    ("YOLO [1]", "Real-time one-stage object detection", CYAN),
    ("ByteTrack [2]", "Temporal association of detections", MINT),
    ("Depth Anything V2 [3]", "Monocular scene-depth estimation", AMBER),
    ("Navigation reviews [7, 8, 10]", "Assistive travel aids and route guidance", RED),
    ("ToF + haptics [9, 12]", "Local navigation and multimodal feedback", CYAN),
    ("User control [13]", "Preserving user agency in assistance", MINT),
]
for i, (head, body, color) in enumerate(research):
    x = 0.85 + (i % 3) * 4.1
    y = 1.55 + (i // 3) * 2.1
    rect(slide, x, y, 3.55, 1.45, fill=PANEL, line=color, rounded=True)
    add_text(slide, head, x + 0.22, y + 0.25, 3.1, 0.35, size=19, color=color, bold=True)
    add_text(slide, body, x + 0.22, y + 0.78, 3.1, 0.42, size=15, color=WHITE)
rect(slide, 1.2, 5.85, 10.8, 0.62, fill=RGBColor(30, 43, 39), line=MINT, rounded=True)
add_text(slide, "Research gap addressed: combine perception, geometry, temporal evidence, risk, and human-centred feedback in one local prototype.", 1.45, 6.03, 10.3, 0.25, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 7. System architecture
slide = base("System architecture", "Five-layer pipeline", 7)
layers = [
    ("INPUT", "IP webcam\nUSB webcam\nvideo clip", CYAN),
    ("PERCEPTION", "YOLO\nByteTrack\nDepth Anything V2", MINT),
    ("REASONING", "Ground plane\n3D points\nforward corridor", AMBER),
    ("DECISION", "Risk score\nTTC\npriority + cooldown", RED),
    ("INTERACTION", "Browser console\nvoice\ncaptions", CYAN),
]
for i, (head, body, color) in enumerate(layers):
    x = 0.65 + i * 2.52
    rect(slide, x, 2.0, 2.05, 2.25, fill=PANEL, line=color, rounded=True)
    add_text(slide, head, x + 0.12, 2.28, 1.8, 0.24, size=11, color=color, bold=True, font="Aptos Mono", align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.18, 2.9, 1.68, 0.9, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if i < len(layers) - 1:
        arrow(slide, x + 2.08, 3.12, x + 2.43, 3.12)
add_text(slide, "Urgent camera obstruction takes priority over routine route or status speech.", 2.05, 5.35, 9.2, 0.35, size=19, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

# 8. System design
slide = base("System design and workflow", "How one frame becomes guidance", 8)
steps = [
    ("01", "Capture", "Read camera or video frame"),
    ("02", "Detect", "Find named obstacles"),
    ("03", "Track", "Persist identities over time"),
    ("04", "Project", "Estimate ground and distance"),
    ("05", "Evaluate", "Test one-metre corridor"),
    ("06", "Speak", "Rank and announce action"),
]
for i, (num, head, body) in enumerate(steps):
    x = 0.7 + (i % 3) * 4.15
    y = 1.55 + (i // 3) * 2.25
    rect(slide, x, y, 3.55, 1.45, fill=PANEL, line=[CYAN, MINT, AMBER, RED, CYAN, MINT][i], rounded=True)
    add_text(slide, num, x + 0.2, y + 0.2, 0.55, 0.28, size=13, color=MINT, bold=True, font="Aptos Mono")
    add_text(slide, head, x + 0.9, y + 0.17, 2.3, 0.32, size=21, color=WHITE, bold=True)
    add_text(slide, body, x + 0.9, y + 0.7, 2.35, 0.35, size=15, color=MUTED)
rect(slide, 1.2, 6.0, 10.8, 0.45, fill=RGBColor(47, 35, 35), line=RED, rounded=True)
add_text(slide, "Example: Stop. Path blocked, 2 meters. Keep left.", 1.45, 6.1, 10.3, 0.22, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 9. Algorithm snippets
slide = base("Algorithm code snippets", "Core decision logic", 9)
code_block(slide, "depth = estimate_depth(frame)\nplane = fit_ground_plane(depth, camera_height)\npoints = project_to_3d(depth, plane)\nblocked = corridor_has_obstacle(points,\n                                width=1.0,\n                                max_range=3.0)", 0.75, 1.45, 5.75, 4.9, "Spatial corridor")
code_block(slide, "tracks = bytetrack.update(detections)\nttc = estimate_ttc(track_history)\npriority = rank_alert(\n    in_path, distance, ttc, persistence\n)\nif priority == 'urgent':\n    speak('Stop. Path blocked. Keep left.')", 6.8, 1.45, 5.75, 4.9, "Risk and alert selection")

# 10. Implementation and requirements
slide = base("Implementation and evaluation", "Prototype scope", 10)
rect(slide, 0.8, 1.5, 5.7, 4.95, fill=PANEL, line=CYAN, rounded=True)
add_text(slide, "Software", 1.2, 1.9, 2.0, 0.35, size=24, color=CYAN, bold=True)
add_bullets(slide, [
    "Python 3.10+ and PyTorch",
    "OpenCV and Ultralytics YOLO11s",
    "ByteTrack and Depth Anything V2",
    "HTML, CSS, JavaScript browser console",
    "Browser speech synthesis and local audio",
], 1.2, 2.65, 4.7, 2.55, size=17)
rect(slide, 6.9, 1.5, 5.55, 4.95, fill=PANEL, line=AMBER, rounded=True)
add_text(slide, "Evaluation", 7.3, 1.9, 2.2, 0.35, size=24, color=AMBER, bold=True)
add_bullets(slide, [
    "Precision, recall, and mAP",
    "Path-blocking accuracy",
    "Distance-estimation error",
    "False-alert and missed-hazard rate",
    "Alert latency and frames per second",
], 7.3, 2.65, 4.5, 2.55, size=17)
add_text(slide, "Testing must use supervision and conventional mobility safety methods.", 1.35, 6.65, 10.7, 0.28, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)

# 11. Novelty
slide = base("Novelty and contribution", "What SafeStride adds", 11)
novelty = [
    ("01", "From boxes to path decisions", "The system asks whether visual evidence is relevant to the forward walking corridor."),
    ("02", "Named + unknown blockers", "YOLO detections are combined with depth-based geometric blocker analysis."),
    ("03", "Approximate distance in speech", "Ground-plane calibration turns relative depth into a practical assistance cue."),
    ("04", "Safety-first interaction", "Urgency, persistence, cooldowns, and route priority shape what the user hears."),
]
for i, (number, head, body) in enumerate(novelty):
    y = 1.45 + i * 1.2
    add_text(slide, number, 0.95, y, 0.55, 0.3, size=14, color=MINT, bold=True, font="Aptos Mono")
    add_text(slide, head, 1.75, y - 0.03, 4.0, 0.35, size=20, color=WHITE, bold=True)
    add_text(slide, body, 5.6, y, 6.2, 0.42, size=17, color=MUTED)
    rect(slide, 0.9, y + 0.68, 11.5, 0.02, fill=LINE)
rect(slide, 2.0, 6.15, 9.2, 0.55, fill=RGBColor(30, 43, 39), line=MINT, rounded=True)
add_text(slide, "Contribution: an integrated, local, educational prototype for distance-aware safe-walking guidance.", 2.2, 6.32, 8.8, 0.22, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# 12. Conclusion
slide = base("Conclusion", "What the project demonstrates", 12)
add_text(slide, "SafeStride Assist connects visual perception to an understandable walking action.", 0.9, 1.55, 11.4, 0.5, size=27, color=MINT, bold=True, align=PP_ALIGN.CENTER)
add_bullets(slide, [
    "YOLO and ByteTrack provide named and temporally stable detections.",
    "Depth Anything V2 and ground-plane reasoning provide approximate distance.",
    "A forward corridor identifies path relevance and unknown blockers.",
    "Risk ranking and cooldowns keep voice feedback concise.",
    "The result is a research prototype, not a certified mobility or emergency device.",
], 1.6, 2.55, 10.0, 2.8, size=20, gap=12)
rect(slide, 2.2, 5.75, 8.9, 0.72, fill=PANEL, line=CYAN, rounded=True)
add_text(slide, "See  ->  reason  ->  prioritise  ->  speak", 2.4, 5.98, 8.5, 0.28, size=22, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

# 13. References I
slide = base("References", "Research and technical sources", 13)
refs_one = [
    "[1] Redmon et al., You Only Look Once, CVPR, 2016.",
    "[2] Zhang et al., ByteTrack, ECCV, 2022.",
    "[3] Yang et al., Depth Anything V2, NeurIPS, 2024.",
    "[4] Ranftl et al., Vision Transformers for Dense Prediction, ICCV, 2021.",
    "[5] Ultralytics, Ultralytics YOLO Documentation.",
    "[6] World Health Organization, World Report on Vision, 2019.",
    "[7] Khan et al., Analysis of Navigation Assistants, IEEE Access, 2021.",
]
add_bullets(slide, refs_one, 0.95, 1.45, 11.3, 4.9, size=17, gap=14)
add_text(slide, "DOIs and URLs are provided in the written synopsis reference list.", 1.0, 6.35, 11.2, 0.3, size=15, color=MUTED, align=PP_ALIGN.CENTER)

# 14. References II
slide = base("References", "Assistive navigation literature", 14)
refs_two = [
    "[8] Xu et al., Wearable Obstacle Avoidance Electronic Travel Aids, IEEE Access, 2023.",
    "[9] Katzschmann et al., Safe Local Navigation with ToF and Haptic Feedback, IEEE TNSRE, 2018.",
    "[10] Eltaher et al., Systematic Review of Urban Navigation Systems, Sensors, 2021.",
    "[11] Johannesson et al., The Sound of Vision Project, Brain Sciences, 2016.",
    "[12] Shull and Damian, Haptic Wearables Review, JNER, 2015.",
    "[13] Ranganeni et al., Levels of Control for Blind Travelers, HRI, 2023.",
    "[14] Zhang et al., ARCore-Based Assistive Navigation, Applied Sciences, 2019.",
]
add_bullets(slide, refs_two, 0.95, 1.45, 11.3, 4.9, size=17, gap=14)
add_text(slide, "Thank you", 1.0, 6.35, 11.2, 0.3, size=18, color=MINT, bold=True, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
